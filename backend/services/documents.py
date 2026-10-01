import difflib
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from uuid import uuid4

import numpy as np

from ..db import SessionFactory
from ..logging_setup import get_logger
from ..models import DocumentChunkORM, DocumentORM
from .embeddings import EmbeddingUnavailable, embed

logger = get_logger("documents")

# Minimum cosine similarity for a chunk to count as relevant on semantic grounds alone.
# Keeps genuinely unrelated questions from dragging in a "closest but still irrelevant"
# chunk just because it's the least-bad option; lexical overlap can still qualify a
# chunk below this bar (e.g. an exact date/ID match embeddings tend to underweight).
_MIN_SEMANTIC_SCORE = 0.3

# Common function words filtered out of lexical matching so two sentences that only
# share "the"/"of"/"is" don't register as topically related — this matters more
# whenever the embedding model isn't available (see EmbeddingUnavailable) and lexical
# overlap is the only relevance signal left.
_STOPWORDS = frozenset(
    "a an the is are was were be been being of in on at to for and or but if so as by "
    "with from that this these those it its i you he she we they what which who whom "
    "when where why how do does did not no yes can could will would should shall may "
    "might must have has had el la los las de en que y por con un una son"
    .split()
)

# Minimum word length before fuzzy matching applies (short words like "is"/"to" would
# false-positive-match almost anything at a reasonable similarity ratio) and the
# similarity ratio a query word needs against a chunk word to count as the same word.
_FUZZY_MIN_WORD_LENGTH = 5
_FUZZY_MATCH_RATIO = 0.82


def _fuzzy_lexical_overlap(query_tokens: set[str], chunk_tokens: set[str]) -> int:
    """Exact overlap first (cheap, common case); falls back to a typo-tolerant fuzzy
    match so a single misspelled letter in a key term (e.g. "restuarant") doesn't drop
    a query below the relevance threshold the way an exact-match-only check would --
    observed happening for real short, informal, code-mixed questions."""
    exact = query_tokens & chunk_tokens
    overlap = len(exact)
    remaining_query = query_tokens - exact
    remaining_chunk = chunk_tokens - exact
    for query_word in remaining_query:
        if len(query_word) < _FUZZY_MIN_WORD_LENGTH:
            continue
        for chunk_word in remaining_chunk:
            if len(chunk_word) < _FUZZY_MIN_WORD_LENGTH:
                continue
            if difflib.SequenceMatcher(None, query_word, chunk_word).ratio() >= _FUZZY_MATCH_RATIO:
                overlap += 1
                break
    return overlap


@dataclass
class StoredDocument:
    id: str
    filename: str
    chunks: list[str] = field(default_factory=list)
    chunk_tokens: list[set[str]] = field(default_factory=list)
    embeddings: np.ndarray | None = None  # shape (len(chunks), dim); None if embedding wasn't available


class DocumentRAG:
    """Local retrieval over uploaded documents: semantic similarity (a local
    multilingual embedding model) as the primary signal, with lexical overlap
    as a small tie-breaker for exact terms embeddings can underweight (dates,
    IDs, names). Persisted to SQLite so uploads survive a server restart.
    """

    def __init__(self, session_factory=SessionFactory) -> None:
        self._session_factory = session_factory
        self.documents: list[StoredDocument] = []
        self._load_from_db()

    def _load_from_db(self) -> None:
        with self._session_factory() as session:
            rows = session.query(DocumentORM).all()
            for row in rows:
                chunks = sorted(row.chunks, key=lambda chunk: chunk.chunk_index)
                texts = [chunk.text for chunk in chunks]
                embeddings = self._stack_stored_embeddings(chunks)
                self.documents.append(
                    StoredDocument(
                        id=row.id,
                        filename=row.filename,
                        chunks=texts,
                        chunk_tokens=[set(self._tokenize(text)) for text in texts],
                        embeddings=embeddings,
                    )
                )

    @staticmethod
    def _stack_stored_embeddings(chunks: list[DocumentChunkORM]) -> np.ndarray | None:
        if not chunks or any(chunk.embedding is None for chunk in chunks):
            return None
        vectors = [np.frombuffer(chunk.embedding, dtype="float32") for chunk in chunks]
        return np.stack(vectors)

    def add_document(self, filename: str, payload: bytes) -> dict[str, object]:
        text = self._extract_text(filename, payload)
        if not text.strip():
            raise ValueError("The uploaded file does not contain readable text.")

        chunks = [self._normalize(chunk) for chunk in self._chunk_text(text)]
        chunk_tokens = [set(self._tokenize(chunk)) for chunk in chunks]
        embeddings = self._try_embed(chunks)

        document_id = str(uuid4())
        with self._session_factory() as session:
            row = DocumentORM(id=document_id, filename=filename)
            session.add(row)
            for index, chunk_text in enumerate(chunks):
                vector_bytes = embeddings[index].tobytes() if embeddings is not None else None
                session.add(
                    DocumentChunkORM(
                        document_id=document_id,
                        chunk_index=index,
                        text=chunk_text,
                        embedding=vector_bytes,
                        embedding_dim=embeddings.shape[1] if embeddings is not None else 0,
                    )
                )
            session.commit()

        self.documents.append(StoredDocument(id=document_id, filename=filename, chunks=chunks, chunk_tokens=chunk_tokens, embeddings=embeddings))
        return {"filename": filename, "chunks": len(chunks), "characters": len(text)}

    def get_context(self, query: str, max_chunks: int = 3) -> str:
        if not self.documents or not query.strip():
            return ""

        normalized_query = self._normalize(query)
        if not normalized_query:
            return ""

        query_tokens = set(self._tokenize(normalized_query))
        query_embedding = self._try_embed([normalized_query])
        query_vector = query_embedding[0] if query_embedding is not None else None

        scored: list[tuple[float, str]] = []
        best_rejected_score = -1.0
        for document in self.documents:
            for index, (chunk_text, chunk_tokens) in enumerate(zip(document.chunks, document.chunk_tokens)):
                lexical_overlap = _fuzzy_lexical_overlap(query_tokens, chunk_tokens)
                semantic_score = 0.0
                if query_vector is not None and document.embeddings is not None:
                    semantic_score = float(np.dot(query_vector, document.embeddings[index]))
                if semantic_score < _MIN_SEMANTIC_SCORE and lexical_overlap == 0:
                    best_rejected_score = max(best_rejected_score, semantic_score)
                    continue
                substring_bonus = 0.1 if normalized_query.lower() in chunk_text.lower() else 0.0
                score = semantic_score + 0.03 * lexical_overlap + substring_bonus
                scored.append((score, chunk_text))

        scored.sort(key=lambda item: item[0], reverse=True)
        top_chunks = [chunk for _, chunk in scored[:max_chunks]]
        if top_chunks:
            logger.info("retrieval hit for %r: %d/%d chunks qualified, top score=%.2f", query, len(scored), sum(len(d.chunks) for d in self.documents), scored[0][0])
        else:
            logger.info(
                "retrieval MISS for %r across %d chunks (closest semantic score=%.2f, threshold=%.2f) -- answer will fall back to 'not enough information'",
                query, sum(len(d.chunks) for d in self.documents), best_rejected_score, _MIN_SEMANTIC_SCORE,
            )
        return "\n\n".join(top_chunks)

    def clear(self) -> None:
        with self._session_factory() as session:
            session.query(DocumentChunkORM).delete()
            session.query(DocumentORM).delete()
            session.commit()
        self.documents.clear()

    def count(self) -> int:
        return len(self.documents)

    @staticmethod
    def _try_embed(texts: list[str]) -> np.ndarray | None:
        try:
            return embed(texts)
        except EmbeddingUnavailable:
            return None

    @staticmethod
    def _normalize(value: str) -> str:
        return re.sub(r"\s+", " ", value or "").strip()

    @staticmethod
    def _tokenize(value: str) -> list[str]:
        tokens = re.findall(r"[A-Za-z0-9][A-Za-z0-9'/-]*", value.lower())
        return [token for token in tokens if token not in _STOPWORDS]

    @staticmethod
    def _chunk_text(text: str, max_chars: int = 900) -> list[str]:
        paragraphs = [part.strip() for part in re.split(r"\n{2,}|\r\n{2,}", text) if part.strip()]
        if not paragraphs:
            return [text[:max_chars]]

        chunks: list[str] = []
        current = ""
        for paragraph in paragraphs:
            if len(current) + len(paragraph) + 2 <= max_chars:
                current = f"{current}\n\n{paragraph}".strip()
            else:
                if current:
                    chunks.append(current)
                current = paragraph
        if current:
            chunks.append(current)
        return chunks or [text[:max_chars]]

    @staticmethod
    def _extract_text(filename: str, payload: bytes) -> str:
        suffix = Path(filename).suffix.lower()

        if suffix in {".txt", ".md", ".csv", ".json", ".yaml", ".yml", ".log"}:
            return payload.decode("utf-8", errors="ignore")

        if suffix == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError as error:
                raise RuntimeError("Install pypdf to enable PDF uploads.") from error

            reader = PdfReader(io.BytesIO(payload))
            return "\n\n".join(page.extract_text() or "" for page in reader.pages)

        if suffix == ".docx":
            try:
                from docx import Document
            except ImportError as error:
                raise RuntimeError("Install python-docx to enable DOCX uploads.") from error

            document = Document(io.BytesIO(payload))
            return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())

        if suffix in {".rtf", ".html", ".htm"}:
            return re.sub(r"<[^>]+>", " ", payload.decode("utf-8", errors="ignore"))

        if suffix == ".doc":
            raise RuntimeError("Legacy .doc files aren't supported — please upload .docx, .pdf, or .txt instead.")

        return payload.decode("utf-8", errors="ignore")


document_rag = DocumentRAG()
